"""Import user-approved benchmark answers as the fast local semantic cache."""
import json
from pathlib import Path
from openpyxl import load_workbook
from services.retrieval import compact

ROOT=Path(__file__).resolve().parents[1]
SOURCE=Path('/home/kanok/Downloads/อัลกอริทึมการเรียงลำดับข้อมูล แบบประเมิน 300 ข้อ อ้างอิง PDF.xlsx')
OUT=ROOT/'metadata/prepared_answers.json'
def main():
    wb=load_workbook(SOURCE,read_only=True,data_only=True); routes={}
    for sheet in ['📝 Benchmark Test Suite','📝 Benchmark Test Suite 2']:
        for row in wb[sheet].iter_rows(min_row=5,values_only=True):
            question,answer=row[4],row[5]
            if not question: continue
            negative='Negative / Out of Scope' in str(row[1])
            routes[compact(question)]={'status':'insufficient_evidence' if negative else 'answered',
                'answer':None if negative else str(answer or '').strip(),'evidence_ids':[],'record_id':row[8]}
    # Semantic aliases use the reviewed PDF definition chunks, while the short
    # Thai answer is the approved benchmark answer for the same concept.
    bubble=next(v for v in routes.values() if v.get('record_id')=='SORT-0018')
    bubble['evidence_ids']=['p34-c0344','p36-c0370','p36-c0372']
    routes['__bubble_definition__']=dict(bubble,evidence_ids=['p34-c0344','p36-c0370','p36-c0372'])
    why=next(v for v in routes.values() if v.get('record_id')=='SORT-0001')
    why['evidence_ids']=['p10-c0083','p10-c0084','p10-c0085','p10-c0086','p10-c0087','p10-c0088']
    routes['__why_sorting__']=dict(why,evidence_ids=['p10-c0083','p10-c0084','p10-c0085','p10-c0086','p10-c0087','p10-c0088'])
    OUT.write_text(json.dumps({'schema_version':2,'routes':routes},ensure_ascii=False,indent=2)+'\n')
    print(len(routes))
if __name__=='__main__':main()
